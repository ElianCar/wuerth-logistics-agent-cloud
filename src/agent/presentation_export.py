from __future__ import annotations

from dataclasses import dataclass, field, replace
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
from typing import Any
from xml.etree import ElementTree
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile, ZipInfo

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx import Presentation
from pptx.util import Inches, Pt

from src.agent.logging_utils import current_timestamp, get_log_dir
from src.agent.presentation_planner import (
    PlannerInvocation,
    PresentationPlanningConfig,
    build_presentation_plan,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE_PATH = PROJECT_ROOT / "assets" / "templates" / "PPT_Vorlage_Wuerth.pptx"
PPTX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
ANTHROPIC_FILES_BETA = "files-api-2025-04-14"
ANTHROPIC_SKILLS_BETAS = (
    "code-execution-2025-08-25",
    "skills-2025-10-02",
    ANTHROPIC_FILES_BETA,
)
ANTHROPIC_PPTX_SKILL_CONTAINER = {
    "skills": [{"type": "anthropic", "skill_id": "pptx", "version": "latest"}]
}
ANTHROPIC_CODE_EXECUTION_TOOL = {"type": "code_execution_20250825", "name": "code_execution"}
DEFAULT_PRESENTATION_MODEL = "claude-opus-4-8"
DEFAULT_PRESENTATION_MAX_TOKENS = 16000
DEFAULT_PRESENTATION_TIMEOUT_SECONDS = 480
DEFAULT_PRESENTATION_MAX_ROWS_FOR_CLAUDE = 80
DEFAULT_PRESENTATION_PAUSE_RETRIES = 8
PRESENTATION_PLANNER_LOG_NAME = "presentation_planner_debug.jsonl"
MAX_TABLE_ROWS_PER_SLIDE = 10
MAX_CHART_POINTS_PER_SLIDE = 50
MAX_TABLE_COLUMNS_PER_SLIDE = 5
MAX_BODY_ITEMS_PER_SLIDE = 8
MAX_BODY_TEXT_CHARS_PER_SLIDE = 700
COVER_SUBTITLE_TEXT_CHARS = 118
EXPECTED_TEMPLATE_SHA256 = "041DE8AC3214DC1892F127021F223D5B9C9D5571B10D6949D022B5A357190EA5"
FIXED_PPTX_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
FIXED_PPTX_CORE_TIMESTAMP = "1980-01-01T00:00:00Z"

LAYOUT_COVER = "agent_01_cover"
LAYOUT_SUMMARY = "agent_02_summary"
LAYOUT_KPI = "agent_03_three_cards"
LAYOUT_COMPARISON = "agent_04_two_column"
LAYOUT_FULL_CONTENT = "agent_05_full_content"
LAYOUT_TWO_CARDS = "agent_06_two_cards"
LAYOUT_CAVEATS = "agent_07_caveats_sources"
LAYOUT_METADATA = "agent_08_appendix_metadata"
LAYOUT_CLOSING = "agent_09_closing"
LAYOUT_CHART = LAYOUT_FULL_CONTENT
LAYOUT_TABLE = LAYOUT_FULL_CONTENT

SLIDE_TYPE_LAYOUTS = {
    "cover": {LAYOUT_COVER},
    "executive_summary": {LAYOUT_SUMMARY},
    "kpi_overview": {LAYOUT_KPI},
    "chart_evidence": {LAYOUT_CHART},
    "table_evidence": {LAYOUT_TABLE},
    "comparison": {LAYOUT_COMPARISON},
    "caveats_sources": {LAYOUT_CAVEATS},
    "appendix_metadata": {LAYOUT_METADATA},
    "closing": {LAYOUT_CLOSING},
}
BODY_REQUIRED_SLIDE_TYPES = {
    "cover",
    "executive_summary",
    "kpi_overview",
    "chart_evidence",
    "comparison",
    "caveats_sources",
    "appendix_metadata",
    "closing",
}
TABLE_REQUIRED_SLIDE_TYPES = {"table_evidence"}
SUPPORTED_CHART_TYPES = {"bar", "line", "top_n_bar"}
DEFAULT_FOOTER_SOURCE = "Wuerth Logistics Analysis"
MAX_EVIDENCE_ROWS = 10
MAX_EVIDENCE_COLUMNS = 5


@dataclass(frozen=True)
class SlideSpec:
    slide_type: str
    layout_name: str
    title: str
    body: list[str] = field(default_factory=list)
    table_columns: list[str] = field(default_factory=list)
    table_rows: list[list[str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, str] = field(default_factory=dict)
    rich_body: list[list[dict[str, Any]]] = field(default_factory=list)


@dataclass(frozen=True)
class SlideDeckSpec:
    title: str
    slides: list[SlideSpec]
    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class TemplateManifest:
    required_layouts: tuple[str, ...]
    optional_layouts: tuple[str, ...] = ()
    layout_aliases: dict[str, str] = field(default_factory=dict)
    placeholder_indexes: dict[str, dict[str, int]] = field(default_factory=dict)
    known_warning_entries: tuple[str, ...] = ()
    expected_sha256: str = ""


@dataclass(frozen=True)
class TemplateAudit:
    available: bool
    template_path: str
    sha256: str = ""
    slide_count: int = 0
    layout_names: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    ole_entries: list[str] = field(default_factory=list)
    external_relationships: list[str] = field(default_factory=list)
    macro_entries: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class _NativeChartPayload:
    chart_type: Any
    data: CategoryChartData


@dataclass(frozen=True)
class PresentationExport:
    available: bool
    content: bytes
    filename: str
    mime_type: str
    slide_count: int
    warnings: list[str]
    unavailable_reason: str = ""
    template_audit: TemplateAudit | None = None
    deck_spec: SlideDeckSpec | None = None


@dataclass(frozen=True)
class PresentationEligibility:
    can_export: bool
    reason: str = ""

    def __bool__(self) -> bool:
        return self.can_export


class PresentationExportError(ValueError):
    """Raised when deterministic PPTX rendering fails unexpectedly."""


class PresentationPlannerInvocationError(RuntimeError):
    """Raised when the optional Sonnet planner request cannot return text."""


class ClaudePresentationExportError(PresentationExportError):
    def __init__(self, reason: str, warning: str) -> None:
        super().__init__(warning)
        self.reason = reason
        self.warning = warning


DEFAULT_TEMPLATE_MANIFEST = TemplateManifest(
    required_layouts=(
        LAYOUT_COVER,
        LAYOUT_SUMMARY,
        LAYOUT_KPI,
        LAYOUT_COMPARISON,
        LAYOUT_FULL_CONTENT,
        LAYOUT_TWO_CARDS,
        LAYOUT_CAVEATS,
        LAYOUT_METADATA,
    ),
    optional_layouts=(LAYOUT_CLOSING,),
    layout_aliases={},
    placeholder_indexes={
        LAYOUT_COVER: {"title": 0, "subtitle": 1, "footer_author": 14, "footer_date": 15},
        LAYOUT_SUMMARY: {"title": 0, "section_label": 13, "body": 14},
        LAYOUT_KPI: {"title": 0, "section_label": 13},
        LAYOUT_COMPARISON: {"title": 0, "section_label": 13, "left_body": 34, "right_body": 32},
        LAYOUT_FULL_CONTENT: {"title": 0, "section_label": 13, "content": 14},
        LAYOUT_TWO_CARDS: {"title": 0, "section_label": 13},
        LAYOUT_CAVEATS: {"title": 0, "section_label": 13, "content": 14},
        LAYOUT_METADATA: {"title": 0, "section_label": 13, "content": 14},
    },
    known_warning_entries=(
        "ppt/embeddings/oleObject1.bin",
        "ppt/embeddings/oleObject2.bin",
    ),
    expected_sha256=EXPECTED_TEMPLATE_SHA256,
)


def build_presentation_export(
    *,
    record: dict[str, Any],
    include_closing: bool = False,
    template_path: Path | str = DEFAULT_TEMPLATE_PATH,
    anthropic_client: Any | None = None,
    planning_config: PresentationPlanningConfig | None = None,
    planner_invocation: PlannerInvocation | None = None,
) -> PresentationExport:
    """Build PPTX bytes from an already validated agent result.

    The default mode uses the local renderer to avoid accidental API spend. Set
    PRESENTATION_EXPORT_MODE=claude explicitly to use Claude's PowerPoint Skill.
    """

    mode = os.getenv("PRESENTATION_EXPORT_MODE", "deterministic").strip().lower()
    if mode in {"deterministic", "local", "python"}:
        return build_deterministic_presentation_export(
            record=record,
            include_closing=include_closing,
            template_path=template_path,
            planning_config=planning_config,
            planner_invocation=planner_invocation,
        )
    if mode in {"claude", "anthropic", "opus"}:
        return build_claude_presentation_export(
            record=record,
            include_closing=include_closing,
            template_path=template_path,
            anthropic_client=anthropic_client,
        )
    return _unavailable_export(
        "presentation_mode_invalid",
        warnings=[f"Unsupported PRESENTATION_EXPORT_MODE: {mode}."],
    )


def build_deterministic_presentation_export(
    *,
    record: dict[str, Any],
    include_closing: bool = False,
    template_path: Path | str = DEFAULT_TEMPLATE_PATH,
    planning_config: PresentationPlanningConfig | None = None,
    planner_invocation: PlannerInvocation | None = None,
) -> PresentationExport:
    """Build deterministic PPTX bytes with the local python-pptx renderer."""

    eligibility = can_export_presentation(record)
    if not eligibility.can_export:
        return _unavailable_export(eligibility.reason)

    try:
        planning_config, planner_invocation = _resolve_planner_dependencies(
            planning_config=planning_config,
            planner_invocation=planner_invocation,
        )
        deck_spec = build_slide_deck_spec(
            record=record,
            include_closing=include_closing,
            planning_config=planning_config,
            planner_invocation=planner_invocation,
        )
    except PresentationExportError as error:
        return _unavailable_export(str(error))

    spec_errors = validate_slide_deck_spec(deck_spec)
    if spec_errors:
        return _unavailable_export("slide_spec_invalid", warnings=spec_errors, deck_spec=deck_spec)

    audit = validate_template(template_path=template_path)
    if audit.errors:
        return _unavailable_export(
            "template_invalid",
            warnings=[*audit.warnings, *audit.errors],
            template_audit=audit,
            deck_spec=deck_spec,
        )

    try:
        content = _render_presentation(
            deck_spec=deck_spec,
            template_path=Path(template_path),
        )
        reopened = Presentation(BytesIO(content))
    except Exception as error:
        return _unavailable_export(
            "render_failed",
            warnings=[f"PPTX rendering failed: {type(error).__name__}."],
            template_audit=audit,
            deck_spec=deck_spec,
        )

    warnings = [*deck_spec.warnings, *audit.warnings]
    return PresentationExport(
        available=True,
        content=content,
        filename=_filename_for(record),
        mime_type=PPTX_MIME_TYPE,
        slide_count=len(reopened.slides),
        warnings=_unique(warnings),
        unavailable_reason="",
        template_audit=audit,
        deck_spec=deck_spec,
    )


def build_sonnet_presentation_planner_invocation(
    config: PresentationPlanningConfig,
    *,
    anthropic_client: Any | None = None,
) -> PlannerInvocation:
    """Build the optional Sonnet JSON-planner invocation for PPT planning."""

    client_holder = anthropic_client

    def invoke(prompt: str) -> str:
        nonlocal client_holder
        log_payload: dict[str, Any] = {
            "timestamp": current_timestamp(),
            "event": "success",
            "mode": config.mode,
            "model": config.model,
            "max_tokens": config.max_tokens,
            "timeout_seconds": config.timeout_seconds,
            "prompt": prompt,
        }
        if client_holder is None:
            try:
                client_holder = _create_anthropic_client()
            except ClaudePresentationExportError as error:
                log_payload.update({"event": "failure", "error_type": type(error).__name__})
                _append_presentation_planner_log(log_payload)
                raise PresentationPlannerInvocationError(error.reason) from error
        try:
            response = _create_presentation_planner_message(
                client_holder,
                prompt=prompt,
                config=config,
            )
            response_text = _response_text(response)
            if not response_text:
                raise PresentationPlannerInvocationError("planner_empty_response")
            log_payload["response_text"] = response_text
            _append_presentation_planner_log(log_payload)
            return response_text
        except Exception as error:
            log_payload.update({"event": "failure", "error_type": type(error).__name__})
            _append_presentation_planner_log(log_payload)
            raise

    return invoke


def _append_presentation_planner_log(payload: dict[str, Any]) -> None:
    try:
        log_path = get_log_dir() / PRESENTATION_PLANNER_LOG_NAME
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with log_path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(payload, ensure_ascii=False, default=str) + "\n")
    except Exception as error:
        print(f"Presentation planner logging failed: {type(error).__name__}")


def _resolve_planner_dependencies(
    *,
    planning_config: PresentationPlanningConfig | None,
    planner_invocation: PlannerInvocation | None,
) -> tuple[PresentationPlanningConfig | None, PlannerInvocation | None]:
    config = planning_config or PresentationPlanningConfig.from_env()
    if config.mode == "llm" and planner_invocation is None:
        return config, build_sonnet_presentation_planner_invocation(config)
    return planning_config, planner_invocation


def build_claude_presentation_export(
    *,
    record: dict[str, Any],
    include_closing: bool = False,
    template_path: Path | str = DEFAULT_TEMPLATE_PATH,
    anthropic_client: Any | None = None,
) -> PresentationExport:
    """Generate a Wuerth PowerPoint deck with Claude's PowerPoint Skill."""

    eligibility = can_export_presentation(record)
    if not eligibility.can_export:
        return _unavailable_export(eligibility.reason)

    audit = _validate_template_for_claude(template_path)
    if audit.errors:
        return _unavailable_export(
            "template_invalid",
            warnings=[*audit.warnings, *audit.errors],
            template_audit=audit,
        )

    try:
        client = anthropic_client or _create_anthropic_client()
    except ClaudePresentationExportError as error:
        return _unavailable_export(
            error.reason,
            warnings=[error.warning],
            template_audit=audit,
        )

    uploaded_file_ids: list[str] = []
    generated_file_id = ""
    try:
        payload = _build_claude_payload(record=record, include_closing=include_closing)
        prompt = _build_claude_presentation_prompt(payload=payload)
        template_file = _upload_claude_file(
            client,
            path=Path(template_path),
            mime_type=PPTX_MIME_TYPE,
        )
        uploaded_file_ids.append(str(template_file.id))
        payload_file = _upload_claude_payload(client, payload)
        uploaded_file_ids.append(str(payload_file.id))

        response = _run_claude_pptx_generation(
            client,
            prompt=prompt,
            template_file_id=str(template_file.id),
            payload_file_id=str(payload_file.id),
        )
        file_ids = _extract_file_ids(response)
        if not file_ids:
            excerpt = _response_text_excerpt(response)
            warning = "Claude did not return a downloadable PPTX file."
            if excerpt:
                warning = f"{warning} Response excerpt: {excerpt}"
            raise ClaudePresentationExportError("claude_output_missing", warning)

        content, generated_file_id, download_warnings = _download_first_pptx(client, file_ids)
        reopened = Presentation(BytesIO(content))
    except ClaudePresentationExportError as error:
        return _unavailable_export(
            error.reason,
            warnings=[*audit.warnings, error.warning],
            template_audit=audit,
        )
    except Exception as error:
        return _unavailable_export(
            "claude_generation_failed",
            warnings=[*audit.warnings, f"Claude PPT generation failed: {type(error).__name__}."],
            template_audit=audit,
        )
    finally:
        _delete_claude_files(client, [*uploaded_file_ids, generated_file_id])

    warnings = _unique([
        *audit.warnings,
        *download_warnings,
        f"Generated with {os.getenv('ANTHROPIC_PRESENTATION_MODEL', DEFAULT_PRESENTATION_MODEL)}.",
    ])
    return PresentationExport(
        available=True,
        content=content,
        filename=_filename_for(record),
        mime_type=PPTX_MIME_TYPE,
        slide_count=len(reopened.slides),
        warnings=warnings,
        unavailable_reason="",
        template_audit=audit,
        deck_spec=None,
    )


def can_export_presentation(record: dict[str, Any]) -> PresentationEligibility:
    """Return whether a record is eligible for PPTX export."""

    reason = _record_unavailable_reason(record)
    return PresentationEligibility(can_export=not reason, reason=reason)


def build_slide_deck_spec(
    *,
    record: dict[str, Any],
    include_closing: bool = False,
    planning_config: PresentationPlanningConfig | None = None,
    planner_invocation: PlannerInvocation | None = None,
) -> SlideDeckSpec:
    """Build an ordered, dynamic slide spec from a successful record."""

    eligibility = can_export_presentation(record)
    if not eligibility.can_export:
        raise PresentationExportError(eligibility.reason)

    reporting = _dict_or_empty(record.get("reporting_result"))
    source_tables = _string_list(record.get("source_tables"))
    plan = build_presentation_plan(
        record=record,
        config=planning_config,
        planner_invocation=planner_invocation,
    )
    title = plan.title
    footer_source = title
    warnings: list[str] = [*plan.warnings, *plan.audit.warnings]

    slides: list[SlideSpec] = [
        SlideSpec(
            slide_type="cover",
            layout_name=LAYOUT_COVER,
            title=title,
            body=[
                _cover_subtitle(record, reporting),
            ],
            metadata={
                "run_id": str(record.get("run_id", "")),
                "footer_author": "Logistics Agent",
                "footer_date": _display_date(record),
                "footer_source": footer_source,
            },
        )
    ]

    summary_body = [bullet.text for bullet in plan.executive_bullets if bullet.text]
    if summary_body:
        slides.append(
            SlideSpec(
                slide_type="executive_summary",
                layout_name=LAYOUT_SUMMARY,
                title="Management-Zusammenfassung",
                body=summary_body,
                metadata={
                    "section_label": "KERNAUSSAGE",
                    "footer_source": footer_source,
                    "footer_date": _display_date(record),
                },
                rich_body=[
                    [{"text": span.text, "bold": span.bold} for span in bullet.spans]
                    for bullet in plan.executive_bullets
                ],
            )
        )

    kpi_cards = _list_of_dicts(reporting.get("kpi_cards"))
    if kpi_cards:
        slides.append(
            SlideSpec(
                slide_type="kpi_overview",
                layout_name=LAYOUT_KPI,
                title="Kennzahlen",
                body=_kpi_body(kpi_cards),
                metadata={
                    **_kpi_card_metadata(kpi_cards),
                    "section_label": "ERGEBNISBILD",
                    "footer_source": footer_source,
                    "footer_date": _display_date(record),
                },
            )
        )

    for chart in plan.charts:
        if chart.render_allowed:
            if not _is_planned_chart_supported(chart):
                warnings.append(f"Geplantes Diagramm '{chart.title}' wurde ausgelassen, weil die Kategorie wie eine Kennzahl wirkt.")
                continue
            chart_rows = [
                [str(row.get("label", "")), str(row.get("value", ""))]
                for row in chart.rows
                if str(row.get("label", "")).strip()
            ]
            if not chart_rows:
                warnings.append("Diagramm wurde ausgelassen, weil keine geplanten Werte vorhanden sind.")
                continue
            body = _chart_evidence_body(chart)
            slides.append(
                SlideSpec(
                    slide_type="chart_evidence",
                    layout_name=LAYOUT_CHART,
                    title=chart.title or "Evidenz",
                    body=body,
                    table_columns=[chart.x_label or "Kategorie", chart.y_label or "Wert"],
                    table_rows=chart_rows,
                    metadata={
                        "chart_type": chart.chart_type,
                        "chart_orientation": chart.orientation,
                        "section_label": "VISUELLE EVIDENZ",
                        "footer_source": footer_source,
                        "footer_date": _display_date(record),
                    },
                )
            )
            continue

        fallback_reason = chart.fallback_reason or "Diagramm wurde nicht gerendert. Die Evidenz wird als Tabelle gezeigt."
        slides.append(
            SlideSpec(
                slide_type="caveats_sources",
                layout_name=LAYOUT_CAVEATS,
                title="Darstellungshinweis",
                body=[
                    fallback_reason,
                    "Die Evidenz wird als geordnete Tabelle gezeigt.",
                ],
                metadata={
                    "section_label": "HINWEIS",
                    "footer_source": footer_source,
                    "footer_date": _display_date(record),
                },
            )
        )

    for page in plan.table_pages:
        table_rows = [
            [str(row.get(column, "")) for column in page.columns]
            for row in page.rows
        ]
        if page.columns and table_rows:
            slides.append(
                SlideSpec(
                    slide_type="table_evidence",
                    layout_name=LAYOUT_TABLE,
                    title="Evidenz",
                    body=page.notes or [page.row_range_label],
                    table_columns=page.columns,
                    table_rows=table_rows,
                    metadata={
                        "section_label": f"TABELLE {page.page_number}",
                        "footer_source": footer_source,
                        "footer_date": _display_date(record),
                    },
                )
            )

    caveat_body = _caveats_body(reporting, source_tables)
    caveat_body = _unique([*plan.caveats, *caveat_body])
    if caveat_body:
        slides.append(
            SlideSpec(
                slide_type="caveats_sources",
                layout_name=LAYOUT_CAVEATS,
                title="Datenbasis und Grenzen",
                body=caveat_body,
                metadata={
                    "section_label": "GRENZEN",
                    "footer_source": footer_source,
                    "footer_date": _display_date(record),
                },
            )
        )

    metadata_body = _metadata_body(record)
    if metadata_body:
        slides.append(
            SlideSpec(
                slide_type="appendix_metadata",
                layout_name=LAYOUT_METADATA,
                title="Technischer Anhang",
                body=metadata_body,
                metadata={
                    "final_sql": str(record.get("final_sql") or ""),
                    "section_label": "TECHNISCHE DETAILS",
                    "footer_source": footer_source,
                    "footer_date": _display_date(record),
                },
            )
        )

    if include_closing:
        slides.append(
            SlideSpec(
                slide_type="closing",
                layout_name=LAYOUT_CLOSING,
                title="Abschluss",
                body=["Erstellt aus einer erfolgreichen und validierten Logistik-Auswertung."],
                metadata={
                    "footer_source": footer_source,
                    "footer_date": _display_date(record),
                },
            )
        )

    return SlideDeckSpec(
        title=title,
        slides=slides,
        warnings=_unique(warnings),
        metadata={
            "run_id": str(record.get("run_id", "")),
            "final_sql": str(record.get("final_sql") or ""),
            "source_tables": ",".join(source_tables),
            "planning_mode": plan.audit.planning_mode,
            "selected_chart_types": ",".join(plan.audit.selected_chart_types),
            "planning_fallback_reasons": " | ".join(plan.audit.fallback_reasons),
            "planner_fallback_reasons": " | ".join(plan.audit.fallback_reasons),
            "planner_warning_codes": ",".join(_planner_warning_codes([*plan.warnings, *plan.audit.warnings])),
        },
    )


def validate_slide_deck_spec(
    deck_spec: SlideDeckSpec,
    *,
    manifest: TemplateManifest = DEFAULT_TEMPLATE_MANIFEST,
) -> list[str]:
    errors: list[str] = []
    allowed_layouts = set(manifest.required_layouts) | set(manifest.optional_layouts)
    repeatable_layouts = {LAYOUT_CHART, LAYOUT_TABLE, LAYOUT_CAVEATS}

    if not deck_spec.slides:
        errors.append("Slide deck spec must contain at least one slide.")
        return errors

    if deck_spec.slides[0].layout_name != LAYOUT_COVER:
        errors.append("First slide must use Agent 01 Cover.")

    layout_names = [slide.layout_name for slide in deck_spec.slides]
    if layout_names == [
        LAYOUT_COVER,
        LAYOUT_SUMMARY,
        LAYOUT_KPI,
        LAYOUT_CHART,
        LAYOUT_TABLE,
        LAYOUT_COMPARISON,
        LAYOUT_CAVEATS,
        LAYOUT_METADATA,
        LAYOUT_CLOSING,
    ]:
        errors.append("Slide deck spec appears to be a fixed 9-slide deck.")

    for index, slide in enumerate(deck_spec.slides, start=1):
        allowed_layouts_for_type = SLIDE_TYPE_LAYOUTS.get(slide.slide_type)
        if allowed_layouts_for_type is None:
            errors.append(f"Slide {index} has unsupported slide type: {slide.slide_type}.")
        elif slide.layout_name not in allowed_layouts_for_type:
            errors.append(
                f"Slide {index} layout {slide.layout_name} is not valid for slide type {slide.slide_type}."
            )
        if slide.layout_name not in allowed_layouts:
            errors.append(f"Slide {index} uses unsupported layout: {slide.layout_name}.")
        if not slide.slide_type:
            errors.append(f"Slide {index} has no slide_type.")
        if not slide.title:
            errors.append(f"Slide {index} has no title.")
        if len(slide.title) > 140:
            errors.append(f"Slide {index} title exceeds text budget.")
        body_text_size = sum(len(item) for item in slide.body)
        if body_text_size > MAX_BODY_TEXT_CHARS_PER_SLIDE:
            errors.append(f"Slide {index} body exceeds text budget.")
        if len(slide.body) > MAX_BODY_ITEMS_PER_SLIDE:
            errors.append(f"Slide {index} body exceeds item budget.")
        if slide.rich_body:
            if len(slide.rich_body) != len(slide.body):
                errors.append(f"Slide {index} rich text run count does not match body item count.")
            for body_index, spans in enumerate(slide.rich_body):
                if body_index >= len(slide.body):
                    continue
                if not isinstance(spans, list) or not spans:
                    errors.append(f"Slide {index} rich text item {body_index + 1} has no runs.")
                    continue
                joined_text = ""
                for span in spans:
                    if not isinstance(span, dict):
                        errors.append(f"Slide {index} rich text item {body_index + 1} has an invalid run.")
                        continue
                    joined_text += str(span.get("text") or "")
                if joined_text != slide.body[body_index]:
                    errors.append(f"Slide {index} rich text item {body_index + 1} does not match body text.")
        if slide.slide_type in BODY_REQUIRED_SLIDE_TYPES and not any(str(item).strip() for item in slide.body):
            errors.append(f"Slide {index} is missing required content.")
        if slide.slide_type in TABLE_REQUIRED_SLIDE_TYPES and (not slide.table_columns or not slide.table_rows):
            errors.append(f"Slide {index} is missing required table content.")
        table_row_limit = MAX_CHART_POINTS_PER_SLIDE if slide.slide_type == "chart_evidence" else MAX_TABLE_ROWS_PER_SLIDE
        if len(slide.table_rows) > table_row_limit:
            errors.append(f"Slide {index} table exceeds row limit.")
        if len(slide.table_columns) > MAX_TABLE_COLUMNS_PER_SLIDE:
            errors.append(f"Slide {index} table exceeds column limit.")
        if slide.table_rows and not slide.table_columns:
            errors.append(f"Slide {index} has rows without table columns.")
        if slide.slide_type == "chart_evidence":
            chart_type = str(slide.metadata.get("chart_type") or "").lower()
            if chart_type not in SUPPORTED_CHART_TYPES:
                errors.append(f"Slide {index} has unsupported chart payload: {chart_type}.")
            if not slide.table_columns or not slide.table_rows:
                errors.append(f"Slide {index} chart evidence is missing required result data.")
            for row_number, row in enumerate(slide.table_rows, start=1):
                if row and len(str(row[0])) > 40:
                    errors.append(f"Slide {index} chart label {row_number} exceeds text budget.")

    repeated = {
        layout_name
        for layout_name in layout_names
        if layout_names.count(layout_name) > 1 and layout_name not in repeatable_layouts
    }
    for layout_name in sorted(repeated):
        errors.append(f"Layout is not repeatable in Phase 1: {layout_name}.")

    return errors


def validate_template(
    *,
    template_path: Path | str = DEFAULT_TEMPLATE_PATH,
    manifest: TemplateManifest = DEFAULT_TEMPLATE_MANIFEST,
) -> TemplateAudit:
    path = Path(template_path)
    safe_path = _safe_path(path)
    if not path.exists():
        return TemplateAudit(
            available=False,
            template_path=safe_path,
            errors=[f"Template file is missing: {safe_path}."],
        )
    if not path.is_file():
        return TemplateAudit(
            available=False,
            template_path=safe_path,
            errors=[f"Template path is not a file: {safe_path}."],
        )

    warnings: list[str] = []
    errors: list[str] = []
    ole_entries: list[str] = []
    macro_entries: list[str] = []
    external_relationships: list[str] = []

    try:
        digest = sha256(path.read_bytes()).hexdigest().upper()
    except OSError as error:
        return TemplateAudit(
            available=False,
            template_path=safe_path,
            errors=[f"Template file is not readable: {type(error).__name__}."],
        )

    if manifest.expected_sha256 and digest != manifest.expected_sha256:
        warnings.append("Template SHA256 differs from the inspected Phase 1 template.")

    try:
        with ZipFile(path) as package:
            names = package.namelist()
            ole_entries = sorted(name for name in names if name.startswith("ppt/embeddings/"))
            macro_entries = sorted(name for name in names if name.lower().endswith("vbaproject.bin"))
            active_x_entries = sorted(name for name in names if "/activeX/" in name or name.startswith("ppt/activeX/"))
            if active_x_entries:
                errors.extend(f"Template contains unsupported ActiveX content: {entry}." for entry in active_x_entries)
            if macro_entries:
                errors.extend(f"Template contains macro content: {entry}." for entry in macro_entries)
            external_relationships = _external_relationships(package)
    except BadZipFile:
        return TemplateAudit(
            available=False,
            template_path=safe_path,
            sha256=digest,
            errors=["Template is not a readable PPTX ZIP package."],
        )
    except OSError as error:
        return TemplateAudit(
            available=False,
            template_path=safe_path,
            sha256=digest,
            errors=[f"Template package scan failed: {type(error).__name__}."],
        )

    if external_relationships:
        errors.extend(
            f"Template contains external relationship: {relationship}."
            for relationship in external_relationships
        )

    if ole_entries:
        warnings.extend(f"Template contains preserved embedded object warning: {entry}." for entry in ole_entries)

    try:
        presentation = Presentation(str(path))
        layout_names = [layout.name for layout in presentation.slide_layouts]
        slide_count = len(presentation.slides)
    except Exception as error:
        return TemplateAudit(
            available=False,
            template_path=safe_path,
            sha256=digest,
            warnings=warnings,
            errors=[*errors, f"Template cannot be opened by python-pptx: {type(error).__name__}."],
            ole_entries=ole_entries,
            external_relationships=external_relationships,
            macro_entries=macro_entries,
        )

    missing_layouts = [layout for layout in manifest.required_layouts if layout not in layout_names]
    if missing_layouts:
        errors.extend(f"Template is missing required layout: {layout}." for layout in missing_layouts)
    if slide_count and presentation.slides[0].slide_layout.name != LAYOUT_COVER:
        errors.append("Template sample slide does not use Agent 01 Cover.")

    return TemplateAudit(
        available=not errors,
        template_path=safe_path,
        sha256=digest,
        slide_count=slide_count,
        layout_names=layout_names,
        warnings=_unique(warnings),
        errors=errors,
        ole_entries=ole_entries,
        external_relationships=external_relationships,
        macro_entries=macro_entries,
    )


def _validate_template_for_claude(template_path: Path | str) -> TemplateAudit:
    audit = validate_template(template_path=template_path)
    allowed_error_fragments = (
        "unexpected embedded object",
        "approved template hash",
    )
    remaining_errors = [
        error
        for error in audit.errors
        if not any(fragment in error.lower() for fragment in allowed_error_fragments)
    ]
    downgraded_errors = [error for error in audit.errors if error not in remaining_errors]
    if not downgraded_errors:
        return audit
    return replace(
        audit,
        available=not remaining_errors,
        errors=remaining_errors,
        warnings=_unique([
            *audit.warnings,
            *downgraded_errors,
            "Embedded template objects were passed through to Claude as template artifacts.",
        ]),
    )


def _create_anthropic_client() -> Any:
    api_key = os.getenv("ANTHROPIC_API_KEY", "").strip()
    if not api_key or api_key == "key":
        raise ClaudePresentationExportError(
            "anthropic_api_key_missing",
            "ANTHROPIC_API_KEY is not configured for Claude PPT generation.",
        )
    try:
        import anthropic
    except ImportError as error:
        raise ClaudePresentationExportError(
            "anthropic_dependency_missing",
            "The anthropic package is not installed. Install requirements.txt before generating PPTX.",
        ) from error
    return anthropic.Anthropic(api_key=api_key)


def _upload_claude_file(client: Any, *, path: Path, mime_type: str) -> Any:
    try:
        with path.open("rb") as handle:
            return client.beta.files.upload(
                file=(path.name, handle, mime_type),
                betas=[ANTHROPIC_FILES_BETA],
                timeout=_presentation_env_int(
                    "ANTHROPIC_PRESENTATION_TIMEOUT_SECONDS",
                    DEFAULT_PRESENTATION_TIMEOUT_SECONDS,
                ),
            )
    except Exception as error:
        raise ClaudePresentationExportError(
            "claude_upload_failed",
            f"Claude file upload failed for {path.name}: {type(error).__name__}.",
        ) from error


def _upload_claude_payload(client: Any, payload: dict[str, Any]) -> Any:
    payload_bytes = json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")
    try:
        return client.beta.files.upload(
            file=("analysis_payload.json", BytesIO(payload_bytes), "application/json"),
            betas=[ANTHROPIC_FILES_BETA],
            timeout=_presentation_env_int(
                "ANTHROPIC_PRESENTATION_TIMEOUT_SECONDS",
                DEFAULT_PRESENTATION_TIMEOUT_SECONDS,
            ),
        )
    except Exception as error:
        raise ClaudePresentationExportError(
            "claude_upload_failed",
            f"Claude analysis payload upload failed: {type(error).__name__}.",
        ) from error


def _run_claude_pptx_generation(
    client: Any,
    *,
    prompt: str,
    template_file_id: str,
    payload_file_id: str,
) -> Any:
    messages: list[dict[str, Any]] = [
        {
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {"type": "container_upload", "file_id": template_file_id},
                {"type": "container_upload", "file_id": payload_file_id},
            ],
        }
    ]
    container: dict[str, Any] = dict(ANTHROPIC_PPTX_SKILL_CONTAINER)
    response = _create_claude_message(client, messages=messages, container=container)

    for _ in range(_presentation_env_int("ANTHROPIC_PRESENTATION_PAUSE_RETRIES", DEFAULT_PRESENTATION_PAUSE_RETRIES)):
        if getattr(response, "stop_reason", "") != "pause_turn":
            return response
        messages.append({"role": "assistant", "content": response.content})
        container_id = getattr(getattr(response, "container", None), "id", "")
        if container_id:
            container = {"id": container_id, **ANTHROPIC_PPTX_SKILL_CONTAINER}
        response = _create_claude_message(client, messages=messages, container=container)

    raise ClaudePresentationExportError(
        "claude_generation_incomplete",
        "Claude PPT generation paused too many times before returning a file.",
    )


def _create_claude_message(client: Any, *, messages: list[dict[str, Any]], container: dict[str, Any]) -> Any:
    model = os.getenv("ANTHROPIC_PRESENTATION_MODEL", DEFAULT_PRESENTATION_MODEL).strip() or DEFAULT_PRESENTATION_MODEL
    max_tokens = _presentation_env_int("ANTHROPIC_PRESENTATION_MAX_TOKENS", DEFAULT_PRESENTATION_MAX_TOKENS)
    timeout = _presentation_env_int("ANTHROPIC_PRESENTATION_TIMEOUT_SECONDS", DEFAULT_PRESENTATION_TIMEOUT_SECONDS)
    try:
        return client.beta.messages.create(
            model=model,
            max_tokens=max_tokens,
            betas=list(ANTHROPIC_SKILLS_BETAS),
            container=container,
            messages=messages,
            tools=[ANTHROPIC_CODE_EXECUTION_TOOL],
            timeout=timeout,
        )
    except Exception as error:
        raise ClaudePresentationExportError(
            "claude_generation_failed",
            f"Claude PPT generation request failed: {type(error).__name__}.",
        ) from error


def _create_presentation_planner_message(
    client: Any,
    *,
    prompt: str,
    config: PresentationPlanningConfig,
) -> Any:
    try:
        return client.messages.create(
            model=config.model,
            max_tokens=config.max_tokens,
            temperature=0,
            messages=[{"role": "user", "content": prompt}],
            timeout=config.timeout_seconds,
        )
    except Exception as error:
        raise PresentationPlannerInvocationError(
            f"planner_request_failed:{type(error).__name__}"
        ) from error


def _download_first_pptx(client: Any, file_ids: list[str]) -> tuple[bytes, str, list[str]]:
    warnings: list[str] = []
    for file_id in file_ids:
        try:
            metadata = client.beta.files.retrieve_metadata(
                file_id=file_id,
                betas=[ANTHROPIC_FILES_BETA],
            )
        except Exception as error:
            warnings.append(f"Could not inspect Claude output file {file_id}: {type(error).__name__}.")
            continue
        filename = str(getattr(metadata, "filename", "") or "")
        mime_type = str(getattr(metadata, "mime_type", "") or "")
        if not filename.lower().endswith(".pptx") and mime_type != PPTX_MIME_TYPE:
            continue
        try:
            content = client.beta.files.download(
                file_id=file_id,
                betas=[ANTHROPIC_FILES_BETA],
            ).read()
            Presentation(BytesIO(content))
            return content, file_id, warnings
        except Exception as error:
            warnings.append(f"Claude output file {filename or file_id} was not a readable PPTX: {type(error).__name__}.")
    raise ClaudePresentationExportError(
        "claude_output_invalid",
        "Claude returned file references, but none was a readable PPTX.",
    )


def _delete_claude_files(client: Any, file_ids: list[str]) -> None:
    for file_id in _unique([file_id for file_id in file_ids if file_id]):
        try:
            client.beta.files.delete(file_id=file_id, betas=[ANTHROPIC_FILES_BETA])
        except Exception:
            continue


def _build_claude_payload(*, record: dict[str, Any], include_closing: bool) -> dict[str, Any]:
    query = _dict_or_empty(record.get("query_result"))
    reporting = _dict_or_empty(record.get("reporting_result"))
    source_tables = _string_list(record.get("source_tables"))
    columns = [str(column) for column in query.get("columns", []) or []]
    rows = _normalized_result_rows(query, columns)
    return {
        "task": "Create a top-management Wuerth PowerPoint from a validated logistics analysis result.",
        "include_closing": bool(include_closing),
        "run": {
            "run_id": str(record.get("run_id") or ""),
            "generated_at": str(record.get("generated_at") or ""),
            "user_question": str(record.get("user_question") or ""),
            "final_answer": str(record.get("final_answer") or ""),
            "final_sql": str(record.get("final_sql") or record.get("generated_sql") or ""),
            "source_tables": source_tables,
            "execution_success": bool(record.get("execution_success")),
            "validation_success": bool(record.get("validation_success") or record.get("sql_valid")),
            "row_count": _effective_row_count(record, query, list(query.get("rows", []) or [])),
        },
        "reporting": {
            "summary": str(reporting.get("summary") or ""),
            "interpretation": str(reporting.get("interpretation") or ""),
            "kpi_cards": _list_of_dicts(reporting.get("kpi_cards")),
            "caveats": [str(item) for item in reporting.get("caveats", []) or [] if str(item).strip()],
            "display_notes": [str(item) for item in reporting.get("display_notes", []) or [] if str(item).strip()],
            "chart_plan": _dict_or_empty(reporting.get("chart_plan") or record.get("chart_spec")),
        },
        "result": {
            "columns": columns,
            "rows_in_payload": len(rows),
            "max_rows_in_payload": _presentation_env_int(
                "PRESENTATION_MAX_ROWS_FOR_CLAUDE",
                DEFAULT_PRESENTATION_MAX_ROWS_FOR_CLAUDE,
            ),
            "rows": rows,
            "profile": _result_profile(columns, rows),
        },
        "deck_rules": {
            "target_audience": "top management",
            "language": "same language as the user question",
            "slide_count_target": "4 to 6 slides",
            "avoid": [
                "raw row dump",
                "one slide per table chunk",
                "unresolved placeholders",
                "overlapping text",
                "off-slide tables",
                "invented facts outside the payload",
            ],
        },
    }


def _build_claude_presentation_prompt(*, payload: dict[str, Any]) -> str:
    question = payload["run"]["user_question"] or "Validated logistics analysis"
    row_count = payload["run"]["row_count"]
    rows_in_payload = payload["result"]["rows_in_payload"]
    include_closing = payload["include_closing"]
    return f"""
Create a Wuerth-branded PowerPoint deck from the uploaded files.

Uploaded files:
- PPT_Vorlage_Wuerth.pptx: use this as the source master/template. Preserve Wuerth brand chrome, logo placement, typography hierarchy, footer style, and clean white/red design language.
- analysis_payload.json: read this file first. It contains the validated SQL result and reporting context.

Business question:
{question}

Data boundary:
- The SQL run is validated and successful.
- Returned row count recorded by the app: {row_count}.
- Rows included in analysis_payload.json: {rows_in_payload}.
- Use only the data in analysis_payload.json. Do not invent causes, totals, or business values that are not present.

Deck requirements:
- Create a concise top-management deck with 4 to 6 slides. Include a closing slide only if analysis_payload.json has include_closing=true. Current include_closing={str(include_closing).lower()}.
- Use the uploaded Wuerth PowerPoint as the actual base template, not a blank deck.
- Remove or replace every placeholder string, including "Click to add text", "Click to add subtitle", "Mastertextformat bearbeiten", and similar master placeholders.
- No overlapping text, no clipped titles, no tables extending beyond slide boundaries, no off-slide objects.
- Do not create one table slide per 5 rows. This is unacceptable.
- Do not dump all result rows. Use a compact evidence table with at most 10 high-signal rows, or aggregate duplicate order numbers when useful.
- If the result is an exception list such as shipments without invoices, make the story explicit: answer, scale of visible evidence, sample exceptions, business risk, recommended next checks, and limitations.
- If a chart would not add insight, skip the chart. A clean table plus callouts is better than a forced chart.
- Put technical details such as SQL, source tables, and caveats in a compact sources/caveats slide or appendix, not across many slides.
- Keep wording crisp and executive. Use the language of the business question.
- Save the final deck as a .pptx file and return it as a downloadable file.

Preferred slide structure:
1. Title slide with the business question and one-line takeaway.
2. Executive answer with 2 to 4 bullets and one prominent key number if supported by the payload.
3. Evidence slide with a readable sample or grouped summary of the exception records.
4. Business implication and recommended follow-up checks.
5. Sources, caveats, and data boundary.
6. Optional closing slide only if requested.

Quality check before returning:
- Inspect the generated slide geometry.
- Fix any placeholder, overlap, clipping, overflow, tiny text, or table-width issue before returning.
- The result should be suitable for top-management review without manual cleanup.
""".strip()


def _normalized_result_rows(query: dict[str, Any], columns: list[str]) -> list[dict[str, str]]:
    max_rows = _presentation_env_int("PRESENTATION_MAX_ROWS_FOR_CLAUDE", DEFAULT_PRESENTATION_MAX_ROWS_FOR_CLAUDE)
    rows = []
    for row in list(query.get("rows", []) or [])[:max_rows]:
        normalized_row: dict[str, str] = {}
        for index, column in enumerate(columns):
            normalized_row[column] = _trim_text(_stringify(_row_value(row, column, index)), 160)
        rows.append(normalized_row)
    return rows


def _result_profile(columns: list[str], rows: list[dict[str, str]]) -> dict[str, Any]:
    profile: dict[str, Any] = {"column_count": len(columns), "row_count_in_payload": len(rows)}
    for preferred_column in ("order_number", "shiptoparty", "customer_material"):
        if preferred_column not in columns:
            continue
        values = [row.get(preferred_column, "") for row in rows if row.get(preferred_column, "")]
        profile[f"distinct_{preferred_column}"] = len(set(values))
        profile[f"top_{preferred_column}_values"] = _top_values(values, limit=8)
    return profile


def _top_values(values: list[str], *, limit: int) -> list[dict[str, Any]]:
    counts: dict[str, int] = {}
    for value in values:
        counts[value] = counts.get(value, 0) + 1
    return [
        {"value": value, "count": count}
        for value, count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))[:limit]
    ]


def _extract_file_ids(value: Any) -> list[str]:
    file_ids: list[str] = []

    def visit(item: Any) -> None:
        if item is None:
            return
        if isinstance(item, dict):
            for key, nested in item.items():
                if key == "file_id" and isinstance(nested, str):
                    file_ids.append(nested)
                else:
                    visit(nested)
            return
        if isinstance(item, (list, tuple)):
            for nested in item:
                visit(nested)
            return
        file_id = getattr(item, "file_id", None)
        if isinstance(file_id, str):
            file_ids.append(file_id)
        for attr in ("content", "output", "outputs", "files"):
            if hasattr(item, attr):
                visit(getattr(item, attr))

    visit(getattr(value, "content", value))
    return _unique(file_ids)


def _response_text_excerpt(response: Any) -> str:
    return _trim_text(" ".join(_response_text(response).split()), 350)


def _response_text(response: Any) -> str:
    chunks: list[str] = []

    def visit(item: Any) -> None:
        if item is None:
            return
        if isinstance(item, dict):
            text = item.get("text")
            if isinstance(text, str):
                chunks.append(text)
            for nested in item.values():
                visit(nested)
            return
        if isinstance(item, (list, tuple)):
            for nested in item:
                visit(nested)
            return
        text = getattr(item, "text", None)
        if isinstance(text, str):
            chunks.append(text)
        if hasattr(item, "content"):
            visit(getattr(item, "content"))

    visit(getattr(response, "content", response))
    return "\n".join(chunk.strip() for chunk in chunks if chunk.strip()).strip()


def _presentation_env_int(name: str, default: int) -> int:
    value = os.getenv(name)
    if value is None:
        return default
    try:
        parsed = int(value)
    except ValueError:
        return default
    return parsed if parsed > 0 else default


def _render_presentation(*, deck_spec: SlideDeckSpec, template_path: Path) -> bytes:
    presentation = Presentation(str(template_path))
    if not presentation.slides:
        cover_slide = presentation.slides.add_slide(_layout_by_name(presentation, LAYOUT_COVER))
    else:
        cover_slide = presentation.slides[0]

    _render_slide(cover_slide, deck_spec.slides[0], slide_number=1, deck_title=deck_spec.title)
    for slide_number, slide_spec in enumerate(deck_spec.slides[1:], start=2):
        slide = presentation.slides.add_slide(_layout_by_name(presentation, slide_spec.layout_name))
        _render_slide(slide, slide_spec, slide_number=slide_number, deck_title=deck_spec.title)

    output = BytesIO()
    presentation.save(output)
    return _normalize_pptx_package(output.getvalue())


def _normalize_pptx_package(content: bytes) -> bytes:
    return _normalize_zip_package(content, normalize_embedded_workbooks=True)


def _normalize_zip_package(content: bytes, *, normalize_embedded_workbooks: bool = False) -> bytes:
    normalized = BytesIO()
    with ZipFile(BytesIO(content), "r") as source, ZipFile(normalized, "w") as target:
        for name in sorted(source.namelist()):
            original = source.getinfo(name)
            info = ZipInfo(filename=name, date_time=FIXED_PPTX_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.comment = original.comment
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = _normalize_core_properties(payload)
            if name == "ppt/viewProps.xml":
                payload = _force_normal_powerpoint_view(payload)
            if normalize_embedded_workbooks and name.endswith(".xlsx"):
                payload = _normalize_zip_package(payload)
            target.writestr(info, payload)
    return normalized.getvalue()


def _normalize_core_properties(payload: bytes) -> bytes:
    timestamp = FIXED_PPTX_CORE_TIMESTAMP.encode("utf-8")
    for tag in (b"created", b"modified"):
        pattern = rb"(<dcterms:" + tag + rb"[^>]*>)(.*?)(</dcterms:" + tag + rb">)"
        payload = re.sub(pattern, lambda match: match.group(1) + timestamp + match.group(3), payload)
    return payload


def _force_normal_powerpoint_view(payload: bytes) -> bytes:
    if b"lastView=" in payload:
        return re.sub(rb'lastView="[^"]+"', b'lastView="sldView"', payload, count=1)
    return re.sub(rb"(<p:viewPr\b)", rb'\1 lastView="sldView"', payload, count=1)


def _render_slide(slide: Any, slide_spec: SlideSpec, *, slide_number: int, deck_title: str) -> None:
    is_cover_slide = slide_spec.layout_name == LAYOUT_COVER
    if slide_spec.layout_name == LAYOUT_COVER:
        _render_cover_slide(slide, slide_spec)
    elif slide_spec.layout_name == LAYOUT_KPI:
        _render_three_card_slide(slide, slide_spec)
    elif slide_spec.layout_name == LAYOUT_TWO_CARDS:
        _render_two_card_slide(slide, slide_spec)
    elif slide_spec.layout_name == LAYOUT_COMPARISON:
        _render_two_column_slide(slide, slide_spec)
    else:
        _render_content_slide(slide, slide_spec)
    if not is_cover_slide:
        _render_common_footer(slide, slide_spec, slide_number=slide_number, deck_title=deck_title)
    _clear_unresolved_placeholder_text(slide)


def _layout_by_name(presentation: Any, layout_name: str) -> Any:
    layout = presentation.slide_layouts.get_by_name(layout_name)
    if layout is None:
        raise PresentationExportError(f"Template layout missing: {layout_name}.")
    return layout


def _render_cover_slide(slide: Any, slide_spec: SlideSpec) -> None:
    cover_text_color = RGBColor(255, 255, 255)
    _set_named_text(slide, "title", slide_spec.title, font_size=48, bold=True, color=cover_text_color)
    subtitle = slide_spec.body[0] if slide_spec.body else ""
    _set_named_text(slide, "subtitle", subtitle, font_size=15, bold=True, color=cover_text_color)
    _set_named_text(
        slide,
        "footer_author",
        slide_spec.metadata.get("footer_author", ""),
        font_size=14,
        color=cover_text_color,
    )
    _set_named_text(
        slide,
        "footer_date",
        slide_spec.metadata.get("footer_date", ""),
        font_size=14,
        color=cover_text_color,
    )


def _render_content_slide(slide: Any, slide_spec: SlideSpec) -> None:
    _render_common_header(slide, slide_spec)
    content_shape = _shape_by_name(slide, "content") or _shape_by_name(slide, "body")
    if content_shape is None:
        return
    if slide_spec.rich_body:
        _set_shape_rich_bullets(content_shape, slide_spec.rich_body, font_size=18)
        return
    if slide_spec.slide_type == "chart_evidence":
        if _replace_shape_with_native_chart(slide, content_shape, slide_spec, notes=slide_spec.body):
            return
        chart_image = _chart_image(slide_spec)
        if chart_image is not None:
            _replace_shape_with_picture(slide, content_shape, chart_image, notes=slide_spec.body)
            return
        fallback_body = [
            "Diagramm konnte nicht gerendert werden.",
            "Evidenzwerte werden tabellarisch gezeigt.",
            *slide_spec.body,
        ]
        _set_shape_bullets(content_shape, fallback_body, font_size=16)
        return
    if slide_spec.table_rows:
        _replace_shape_with_table(
            slide,
            content_shape,
            slide_spec.table_columns,
            slide_spec.table_rows,
            notes=slide_spec.body,
        )
        return
    _set_shape_bullets(content_shape, slide_spec.body, font_size=18)


def _render_three_card_slide(slide: Any, slide_spec: SlideSpec) -> None:
    _render_common_header(slide, slide_spec)
    for index in range(1, 4):
        value = slide_spec.metadata.get(f"card_{index}_value", "")
        title = slide_spec.metadata.get(f"card_{index}_title", "")
        body = slide_spec.metadata.get(f"card_{index}_body", "")
        if not value and not title and not body:
            _remove_named_shape(slide, f"card_{index}_visual")
            _remove_named_shape(slide, f"card_{index}_title")
            _remove_named_shape(slide, f"card_{index}_body")
            continue
        _set_named_text(slide, f"card_{index}_visual", value, font_size=30, bold=True, center=True)
        _set_named_text(slide, f"card_{index}_title", title, font_size=18, bold=True)
        _set_named_text(slide, f"card_{index}_body", body, font_size=13)


def _render_two_card_slide(slide: Any, slide_spec: SlideSpec) -> None:
    _render_common_header(slide, slide_spec)
    for index in range(1, 3):
        _set_named_text(
            slide,
            f"card_{index}_title",
            slide_spec.metadata.get(f"card_{index}_title", ""),
            font_size=18,
            bold=True,
        )
        content_shape = _shape_by_name(slide, f"card_{index}_content")
        if content_shape is not None:
            body = slide_spec.metadata.get(f"card_{index}_body", "")
            _set_shape_bullets(content_shape, [body] if body else [], font_size=15)


def _render_two_column_slide(slide: Any, slide_spec: SlideSpec) -> None:
    _render_common_header(slide, slide_spec)
    midpoint = max(1, (len(slide_spec.body) + 1) // 2)
    left_body = _string_list_from_any(slide_spec.body[:midpoint])
    right_body = _string_list_from_any(slide_spec.body[midpoint:])
    _set_named_bullets(slide, "left_body", left_body, font_size=16)
    _set_named_bullets(slide, "right_body", right_body or ["No additional notes."], font_size=16)


def _render_common_header(slide: Any, slide_spec: SlideSpec) -> None:
    _set_named_text(slide, "title", slide_spec.title, font_size=22, bold=True, color=RGBColor(0, 0, 0))
    _set_named_text(
        slide,
        "section_label",
        slide_spec.metadata.get("section_label", ""),
        font_size=18,
        bold=True,
        color=RGBColor(210, 0, 0),
    )


def _render_common_footer(slide: Any, slide_spec: SlideSpec, *, slide_number: int, deck_title: str) -> None:
    footer_source = slide_spec.metadata.get("footer_source") or deck_title or DEFAULT_FOOTER_SOURCE
    _set_named_text(slide, "footer_source", footer_source, font_size=7, color=RGBColor(0, 0, 0))
    _set_named_text(slide, "footer_date", slide_spec.metadata.get("footer_date", ""), font_size=7, color=RGBColor(0, 0, 0))
    _set_named_text(slide, "footer_slide_number", str(slide_number), font_size=7, color=RGBColor(0, 0, 0))


def _shape_by_name(slide: Any, name: str) -> Any | None:
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    layout_shape = _layout_shape_by_name(slide, name)
    if layout_shape is None or not getattr(layout_shape, "is_placeholder", False):
        return None
    placeholder_idx = layout_shape.placeholder_format.idx
    for shape in slide.shapes:
        if not getattr(shape, "is_placeholder", False):
            continue
        if shape.placeholder_format.idx == placeholder_idx:
            return shape
    return None


def _layout_shape_by_name(slide: Any, name: str) -> Any | None:
    layout = getattr(slide, "slide_layout", None)
    if layout is None:
        return None
    for shape in layout.shapes:
        if shape.name == name:
            return shape
    return None


def _add_layout_overlay_textbox(slide: Any, name: str) -> Any | None:
    layout_shape = _layout_shape_by_name(slide, name)
    if layout_shape is None:
        return None
    shape = slide.shapes.add_textbox(layout_shape.left, layout_shape.top, layout_shape.width, layout_shape.height)
    _copy_text_frame_layout(layout_shape, shape)
    shape.fill.background()
    shape.line.fill.background()
    shape.name = name
    return shape


def _copy_text_frame_layout(source_shape: Any, target_shape: Any) -> None:
    if not getattr(source_shape, "has_text_frame", False) or not getattr(target_shape, "has_text_frame", False):
        return
    source_frame = source_shape.text_frame
    target_frame = target_shape.text_frame
    for attribute in (
        "margin_left",
        "margin_right",
        "margin_top",
        "margin_bottom",
        "vertical_anchor",
        "word_wrap",
    ):
        try:
            value = getattr(source_frame, attribute)
            if value is not None:
                setattr(target_frame, attribute, value)
        except (AttributeError, TypeError, ValueError):
            continue


def _set_named_text(
    slide: Any,
    name: str,
    text: str,
    *,
    font_size: int,
    bold: bool = False,
    center: bool = False,
    color: RGBColor | None = None,
) -> None:
    shape = _shape_by_name(slide, name)
    if shape is None and name.startswith("footer_") and str(text).strip():
        shape = _add_layout_overlay_textbox(slide, name)
    if shape is None:
        return
    _set_shape_text(shape, text, font_size=font_size, bold=bold, center=center, color=color)


def _set_named_bullets(slide: Any, name: str, items: list[str], *, font_size: int) -> None:
    shape = _shape_by_name(slide, name)
    if shape is None:
        return
    _set_shape_bullets(shape, items, font_size=font_size)


def _set_shape_text(
    shape: Any,
    text: str,
    *,
    font_size: int,
    bold: bool = False,
    center: bool = False,
    color: RGBColor | None = None,
) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = _trim_text(text, 350)
    paragraph.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    for run in paragraph.runs:
        run.font.size = Pt(font_size)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color


def _set_shape_bullets(shape: Any, items: list[str], *, font_size: int) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    clean_items = [_trim_text(item, 220) for item in items if str(item).strip()]
    if not clean_items:
        return
    for index, item in enumerate(clean_items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        for run in paragraph.runs:
            run.font.size = Pt(font_size)


def _set_shape_rich_bullets(shape: Any, rich_items: list[list[dict[str, Any]]], *, font_size: int) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for item_index, spans in enumerate(rich_items[:MAX_BODY_ITEMS_PER_SLIDE]):
        paragraph = frame.paragraphs[0] if item_index == 0 else frame.add_paragraph()
        paragraph.level = 0
        for span in spans:
            text = str(span.get("text") or "")
            if not text:
                continue
            run = paragraph.add_run()
            run.text = text
            run.font.size = Pt(font_size)
            run.font.bold = bool(span.get("bold"))


def _replace_shape_with_table(
    slide: Any,
    shape: Any,
    columns: list[str],
    rows: list[list[str]],
    *,
    notes: list[str] | None = None,
) -> None:
    bounds = (shape.left, shape.top, shape.width, shape.height)
    _remove_shape(shape)
    note_items = [str(note) for note in notes or [] if str(note).strip()]
    if note_items:
        note_height = Inches(0.48)
        note_shape = slide.shapes.add_textbox(bounds[0], bounds[1], bounds[2], note_height)
        note_shape.name = "table_notes"
        _set_shape_bullets(note_shape, [" | ".join(note_items)], font_size=9)
        table_bounds = (bounds[0], bounds[1] + note_height, bounds[2], max(Inches(0.5), bounds[3] - note_height))
    else:
        table_bounds = bounds
    _add_table(slide, bounds=table_bounds, columns=columns[:MAX_EVIDENCE_COLUMNS], rows=rows[:MAX_EVIDENCE_ROWS])


def _replace_shape_with_picture(
    slide: Any,
    shape: Any,
    image: BytesIO,
    *,
    notes: list[str] | None = None,
) -> None:
    bounds = (shape.left, shape.top, shape.width, shape.height)
    _remove_shape(shape)
    note_items = [str(note) for note in notes or [] if str(note).strip()]
    if note_items:
        note_height = Inches(0.48)
        note_shape = slide.shapes.add_textbox(bounds[0], bounds[1], bounds[2], note_height)
        note_shape.name = "chart_notes"
        _set_shape_bullets(note_shape, [" | ".join(note_items)], font_size=9)
        picture_bounds = (
            bounds[0],
            bounds[1] + note_height,
            bounds[2],
            max(Inches(0.5), bounds[3] - note_height),
        )
    else:
        picture_bounds = bounds
    slide.shapes.add_picture(
        image,
        picture_bounds[0],
        picture_bounds[1],
        width=picture_bounds[2],
        height=picture_bounds[3],
    )


def _replace_shape_with_native_chart(
    slide: Any,
    shape: Any,
    slide_spec: SlideSpec,
    *,
    notes: list[str] | None = None,
) -> bool:
    payload = _native_chart_payload(slide_spec)
    if payload is None:
        return False

    bounds = (shape.left, shape.top, shape.width, shape.height)
    note_items = [str(note) for note in notes or [] if str(note).strip()]
    chart_bounds = _bounds_below_note(bounds, note_items)
    try:
        chart_shape = slide.shapes.add_chart(
            payload.chart_type,
            chart_bounds[0],
            chart_bounds[1],
            chart_bounds[2],
            chart_bounds[3],
            payload.data,
        )
    except Exception:
        return False

    _remove_shape(shape)
    chart_shape.name = "editable_chart"
    _format_native_chart(chart_shape.chart, slide_spec)
    if note_items:
        _add_note_textbox(slide, bounds, name="chart_notes", notes=note_items)
    return True


def _bounds_below_note(bounds: tuple[int, int, int, int], note_items: list[str]) -> tuple[int, int, int, int]:
    if not note_items:
        return bounds
    note_height = Inches(0.48)
    return (
        bounds[0],
        bounds[1] + note_height,
        bounds[2],
        max(Inches(0.5), bounds[3] - note_height),
    )


def _add_note_textbox(slide: Any, bounds: tuple[int, int, int, int], *, name: str, notes: list[str]) -> None:
    note_height = Inches(0.48)
    note_shape = slide.shapes.add_textbox(bounds[0], bounds[1], bounds[2], note_height)
    note_shape.name = name
    _set_shape_bullets(note_shape, [" | ".join(notes)], font_size=9)


def _native_chart_payload(slide_spec: SlideSpec) -> _NativeChartPayload | None:
    if not slide_spec.table_columns or not slide_spec.table_rows:
        return None
    chart_type = str(slide_spec.metadata.get("chart_type", "bar") or "bar").lower()
    native_chart_type = _native_chart_type(slide_spec)
    if native_chart_type is None:
        return None

    categories = [_trim_text_without_ellipsis(row[0], 32) for row in slide_spec.table_rows if row]
    values: list[float] = []
    for row in slide_spec.table_rows:
        try:
            values.append(float(str(row[1]).replace(",", "")))
        except (IndexError, ValueError):
            return None
    if not categories or len(categories) != len(values):
        return None

    data = CategoryChartData()
    if chart_type in {"bar", "top_n_bar"} and slide_spec.metadata.get("chart_orientation") == "horizontal":
        data.categories = list(reversed(categories))
        series_values = list(reversed(values))
    else:
        data.categories = categories
        series_values = values
    y_label = _display_column_name(slide_spec.table_columns[1] if len(slide_spec.table_columns) > 1 else "Wert")
    data.add_series(y_label, series_values)
    return _NativeChartPayload(chart_type=native_chart_type, data=data)


def _native_chart_type(slide_spec: SlideSpec) -> Any | None:
    chart_type = str(slide_spec.metadata.get("chart_type", "bar") or "bar").lower()
    if chart_type == "line":
        return XL_CHART_TYPE.LINE_MARKERS
    if chart_type in {"bar", "top_n_bar"} and slide_spec.metadata.get("chart_orientation") == "horizontal":
        return XL_CHART_TYPE.BAR_CLUSTERED
    if chart_type in {"bar", "top_n_bar"}:
        return XL_CHART_TYPE.COLUMN_CLUSTERED
    return None


def _format_native_chart(chart: Any, slide_spec: SlideSpec) -> None:
    try:
        chart.has_legend = False
        chart.has_title = True
        chart.chart_title.text_frame.text = slide_spec.title
        chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        chart.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    except Exception:
        pass

    try:
        series = chart.series[0]
        if str(slide_spec.metadata.get("chart_type") or "").lower() == "line":
            series.format.line.color.rgb = RGBColor(210, 0, 0)
            series.format.line.width = Pt(2.25)
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor(210, 0, 0)
            series.format.line.color.rgb = RGBColor(210, 0, 0)
    except Exception:
        pass

    for axis_name in ("category_axis", "value_axis"):
        try:
            axis = getattr(chart, axis_name)
            axis.tick_labels.font.size = Pt(8)
            axis.tick_labels.font.color.rgb = RGBColor(0, 0, 0)
        except Exception:
            pass


def _add_table(
    slide: Any,
    *,
    bounds: tuple[int, int, int, int],
    columns: list[str],
    rows: list[list[str]],
) -> None:
    if not columns or not rows:
        return
    left, top, width, height = bounds
    visible_rows = rows[:MAX_EVIDENCE_ROWS]
    shape = slide.shapes.add_table(len(visible_rows) + 1, len(columns), left, top, width, height)
    table = shape.table
    for column_index, column_name in enumerate(columns):
        cell = table.cell(0, column_index)
        cell.text = _display_column_name(column_name)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(210, 0, 0)
        _format_cell_text(cell, font_size=8, bold=True, color=RGBColor(255, 255, 255))
    for row_index, row in enumerate(visible_rows, start=1):
        for column_index, value in enumerate(row[:len(columns)]):
            cell = table.cell(row_index, column_index)
            cell.text = _trim_text(value, 60)
            if row_index % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(242, 222, 222)
            _format_cell_text(cell, font_size=8, bold=False, color=RGBColor(0, 0, 0))


def _format_cell_text(cell: Any, *, font_size: int, bold: bool, color: RGBColor) -> None:
    cell.margin_left = Inches(0.04)
    cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = color


def _chart_image(slide_spec: SlideSpec) -> BytesIO | None:
    if not slide_spec.table_columns or not slide_spec.table_rows:
        return None
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return None

    x_values = [_trim_text(row[0], 32) for row in slide_spec.table_rows]
    y_values: list[float] = []
    for row in slide_spec.table_rows:
        try:
            y_values.append(float(str(row[1]).replace(",", "")))
        except (IndexError, ValueError):
            return None

    fig, ax = plt.subplots(figsize=(10.5, 4.4), dpi=160)
    chart_type = slide_spec.metadata.get("chart_type", "bar")
    if chart_type == "line":
        ax.plot(x_values, y_values, color="#D00000", linewidth=2.5, marker="o")
        ax.set_xlabel(_display_column_name(slide_spec.table_columns[0]), fontsize=9)
        ax.set_ylabel(_display_column_name(slide_spec.table_columns[1]), fontsize=9)
        ax.tick_params(axis="x", labelrotation=30, labelsize=8)
        ax.tick_params(axis="y", labelsize=8)
    elif chart_type in {"bar", "top_n_bar"} and slide_spec.metadata.get("chart_orientation") == "horizontal":
        labels = list(reversed(x_values))
        values = list(reversed(y_values))
        ax.barh(labels, values, color="#D00000")
        ax.set_xlabel(_display_column_name(slide_spec.table_columns[1]), fontsize=9)
        ax.set_ylabel(_display_column_name(slide_spec.table_columns[0]), fontsize=9)
        ax.tick_params(axis="x", labelsize=8)
        ax.tick_params(axis="y", labelsize=8)
    else:
        ax.bar(x_values, y_values, color="#D00000")
        ax.set_xlabel(_display_column_name(slide_spec.table_columns[0]), fontsize=9)
        ax.set_ylabel(_display_column_name(slide_spec.table_columns[1]), fontsize=9)
        ax.tick_params(axis="x", labelrotation=30, labelsize=8)
        ax.tick_params(axis="y", labelsize=8)
    ax.set_title(slide_spec.title, loc="left", fontsize=13, fontweight="bold")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    output = BytesIO()
    fig.savefig(output, format="png", transparent=False, bbox_inches="tight")
    plt.close(fig)
    output.seek(0)
    return output


def _remove_named_shape(slide: Any, name: str) -> None:
    shape = _shape_by_name(slide, name)
    if shape is not None:
        _remove_shape(shape)


def _remove_shape(shape: Any) -> None:
    element = shape._element
    element.getparent().remove(element)


def _clear_unresolved_placeholder_text(slide: Any) -> None:
    placeholder_fragments = (
        "Click to add",
        "Mastertextformat",
        "Mastertitelformat",
        "Master-Untertitelformat",
        "Inhalte einfügen",
        "Schlagwort hinzufügen",
        "Erläuterung hinzufügen",
    )
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs)
        if any(fragment in text for fragment in placeholder_fragments):
            shape.text_frame.clear()


def _record_unavailable_reason(record: dict[str, Any]) -> str:
    if not isinstance(record, dict):
        return "record_missing"
    if record.get("blocked_or_unsafe"):
        return "blocked_request"
    if record.get("needs_clarification"):
        return "clarification_needed"
    if not record.get("execution_success"):
        return "sql_execution_failed"
    if record.get("validation_success") is not True:
        return "sql_validation_failed"
    if record.get("sql_valid") is not True:
        return "sql_validation_failed"

    query = record.get("query_result")
    if not isinstance(query, dict):
        return "missing_query_result"
    columns = query.get("columns") or []
    rows = query.get("rows") or []
    if not columns:
        return "missing_query_columns"
    if not rows:
        return "missing_query_rows"
    effective_row_count = _effective_row_count(record, query, rows)
    if effective_row_count == 0:
        return "zero_row_count"
    return ""


def _unavailable_export(
    reason: str,
    *,
    warnings: list[str] | None = None,
    template_audit: TemplateAudit | None = None,
    deck_spec: SlideDeckSpec | None = None,
) -> PresentationExport:
    return PresentationExport(
        available=False,
        content=b"",
        filename="",
        mime_type=PPTX_MIME_TYPE,
        slide_count=0,
        warnings=warnings or [],
        unavailable_reason=reason,
        template_audit=template_audit,
        deck_spec=deck_spec,
    )


def _summary_body(reporting: dict[str, Any]) -> list[str]:
    body: list[str] = []
    summary = str(reporting.get("summary") or "").strip()
    interpretation = str(reporting.get("interpretation") or "").strip()
    if interpretation:
        body.append(_trim_text(interpretation, 220))
    body.extend(_extract_summary_lines(summary, limit=4))
    return _unique([item for item in body if item])[:MAX_BODY_ITEMS_PER_SLIDE]


def _cover_subtitle(record: dict[str, Any], reporting: dict[str, Any]) -> str:
    interpretation = str(reporting.get("interpretation") or "").strip()
    if interpretation:
        return _trim_text_without_ellipsis(interpretation, COVER_SUBTITLE_TEXT_CHARS)
    final_answer = str(record.get("final_answer") or "").strip()
    if final_answer:
        return _trim_text_without_ellipsis(final_answer, COVER_SUBTITLE_TEXT_CHARS)
    return "Validierte Logistik-Auswertung"


def _display_date(record: dict[str, Any]) -> str:
    generated_at = str(record.get("generated_at") or "").strip()
    for value in (generated_at, _date_from_run_id(record), date.today()):
        formatted = _format_display_date(value)
        if formatted:
            return formatted
    return ""


def _date_from_run_id(record: dict[str, Any]) -> str:
    run_id = str(record.get("run_id") or "").strip()
    match = re.search(r"(?:^|_)((?:19|20)\d{6})(?:_|$)", run_id)
    return match.group(1) if match else ""


def _format_display_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    text = str(value or "").strip()
    if not text:
        return ""
    if re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", text):
        return text
    for pattern, candidate in (
        ("%Y-%m-%d", text[:10]),
        ("%Y%m%d", text[:8]),
    ):
        try:
            return datetime.strptime(candidate, pattern).strftime("%d.%m.%Y")
        except ValueError:
            continue
    return ""


def _kpi_card_metadata(kpi_cards: list[dict[str, Any]]) -> dict[str, str]:
    metadata: dict[str, str] = {}
    for index, card in enumerate(kpi_cards[:3], start=1):
        label = _text_or_default(card.get("label"), "KPI")
        value = _text_or_default(card.get("value"), "")
        note = str(card.get("note") or "").strip()
        metadata[f"card_{index}_title"] = label
        metadata[f"card_{index}_value"] = value
        metadata[f"card_{index}_body"] = _trim_text(note, 120)
    return metadata


def _extract_summary_lines(summary: str, *, limit: int) -> list[str]:
    if not summary:
        return []
    lines = []
    for part in re.split(r"\n+|(?<=\.)\s+", summary):
        text = _trim_text(part.strip(), 220)
        if text:
            lines.append(text)
        if len(lines) >= limit:
            break
    return lines


def _kpi_body(kpi_cards: list[dict[str, Any]]) -> list[str]:
    body = []
    for card in kpi_cards[:3]:
        label = _text_or_default(card.get("label"), "KPI")
        value = _text_or_default(card.get("value"), "")
        note = str(card.get("note") or "").strip()
        body.append(_trim_text(f"{label}: {value}" + (f" ({note})" if note else ""), 180))
    return body


def _chart_body(chart_plan: dict[str, Any], reporting: dict[str, Any]) -> list[str]:
    body = [
        f"Chart type: {_text_or_default(chart_plan.get('chart_type'), 'unknown')}",
        f"X axis: {_text_or_default(chart_plan.get('x_axis'), 'not provided')}",
        f"Y axis: {_text_or_default(chart_plan.get('y_axis'), 'not provided')}",
    ]
    reason = str(chart_plan.get("reason") or "").strip()
    if reason:
        body.append(_trim_text(reason, 200))
    interpretation = str(reporting.get("interpretation") or "").strip()
    if interpretation:
        body.append(_trim_text(interpretation, 220))
    return body[:MAX_BODY_ITEMS_PER_SLIDE]


def _chart_evidence_body(chart: Any) -> list[str]:
    body = [*list(getattr(chart, "notes", []) or []), "Diagramm aus validierten Ergebnisdaten."]
    rows = list(getattr(chart, "rows", []) or [])
    if any(str(row.get("label", "")) == "Sonstige" for row in rows if isinstance(row, dict)):
        body.append("Sonstige buendelt weitere Kategorien.")
    return _unique(body)


def _is_planned_chart_supported(chart: Any) -> bool:
    category_column = str(getattr(chart, "category_column", "") or "").lower()
    chart_type = str(getattr(chart, "chart_type", "") or "").lower()
    if chart_type != "top_n_bar":
        return True
    if category_column in {"order_number", "shiptoparty", "customer_material"}:
        return True
    measure_tokens = ("count", "sum", "total", "amount", "quantity", "rows", "metric", "value")
    return not any(token in category_column for token in measure_tokens)


def _chart_table_columns(chart_plan: dict[str, Any]) -> list[str]:
    columns = [chart_plan.get("x_axis"), chart_plan.get("y_axis")]
    return [str(column) for column in columns if column]


def _chart_table_rows(query: dict[str, Any], chart_plan: dict[str, Any]) -> list[list[str]]:
    columns = [str(column) for column in query.get("columns", []) or []]
    chart_columns = _chart_table_columns(chart_plan)
    if len(chart_columns) != 2 or not all(column in columns for column in chart_columns):
        return []
    indexes = [columns.index(column) for column in chart_columns]
    rows = []
    for row in list(query.get("rows", []) or [])[:MAX_TABLE_ROWS_PER_SLIDE]:
        rows.append([_stringify(_row_value(row, columns[index], index)) for index in indexes])
    return rows


def _table_content(query: dict[str, Any]) -> tuple[list[str], list[list[str]], list[str]]:
    columns = [str(column) for column in query.get("columns", []) or []]
    rows = list(query.get("rows", []) or [])
    warnings: list[str] = []
    if len(columns) > MAX_TABLE_COLUMNS_PER_SLIDE:
        warnings.append("Result table columns were truncated for Phase 1 slide readability.")
        columns = columns[:MAX_TABLE_COLUMNS_PER_SLIDE]
    table_rows = []
    for row in rows:
        table_rows.append([
            _stringify(_row_value(row, column, index))
            for index, column in enumerate(columns)
        ])
    if len(table_rows) > MAX_TABLE_ROWS_PER_SLIDE:
        warnings.append("Result table was sampled for slide readability.")
    return columns, table_rows, warnings


def _caveats_body(reporting: dict[str, Any], source_tables: list[str]) -> list[str]:
    body = [str(caveat) for caveat in reporting.get("caveats", []) or [] if str(caveat).strip()]
    if source_tables:
        body.append(f"Quelltabellen: {_source_text(source_tables)}")
    body.extend(str(note) for note in reporting.get("display_notes", []) or [] if str(note).strip())
    return [_trim_text(item, 220) for item in body[:MAX_BODY_ITEMS_PER_SLIDE]]


def _metadata_body(record: dict[str, Any]) -> list[str]:
    source_tables = _string_list(record.get("source_tables"))
    body = [
        f"Run ID: {_text_or_default(record.get('run_id'), 'nicht erfasst')}",
        f"Ausfuehrung erfolgreich: {bool(record.get('execution_success'))}",
        f"SQL-Validierung: {bool(record.get('validation_success') or record.get('sql_valid'))}",
        f"Zeilen im Ergebnis: {int(record.get('row_count', 0) or 0)}",
        f"Quelltabellen: {_source_text(source_tables)}",
        f"Erstellt am: {_text_or_default(record.get('generated_at'), 'nicht erfasst')}",
    ]
    candidates = record.get("template_candidates")
    if isinstance(candidates, list) and candidates:
        ids = [
            str(candidate.get("template_id") or candidate.get("candidate_id") or candidate.get("id"))
            for candidate in candidates
            if isinstance(candidate, dict)
        ]
        if ids:
            body.append(f"Memory-Template-IDs: {', '.join(ids)}")
    return body[:MAX_BODY_ITEMS_PER_SLIDE]


def _external_relationships(package: ZipFile) -> list[str]:
    relationships = []
    for name in package.namelist():
        if not name.endswith(".rels"):
            continue
        try:
            root = ElementTree.fromstring(package.read(name))
        except (ElementTree.ParseError, KeyError):
            continue
        for relationship in root:
            if relationship.attrib.get("TargetMode", "").lower() == "external":
                relationships.append(name)
                break
    return sorted(relationships)


def _dict_or_empty(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _list_of_dicts(value: Any) -> list[dict[str, Any]]:
    return [item for item in value or [] if isinstance(item, dict)] if isinstance(value, list) else []


def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item) for item in value if str(item)]


def _deck_title(record: dict[str, Any]) -> str:
    question = str(record.get("user_question") or "").strip()
    if question:
        return _trim_text(question, 90)
    return "Wuerth Logistics Analysis"


def _filename_for(record: dict[str, Any]) -> str:
    run_id = str(record.get("run_id") or "analysis")
    safe_run_id = re.sub(r"[^A-Za-z0-9_.-]+", "_", run_id).strip("_") or "analysis"
    return f"wuerth_logistics_{safe_run_id}.pptx"


def _source_text(source_tables: list[str]) -> str:
    return ", ".join(source_tables) if source_tables else "nicht erfasst"


def _display_column_name(column: str) -> str:
    return str(column).replace("_", " ").strip().title()


def _text_or_default(value: Any, default: str) -> str:
    text = str(value or "").strip()
    return text if text else default


def _trim_text(value: str, limit: int) -> str:
    text = re.sub(r"\s+", " ", str(value)).strip()
    if len(text) <= limit:
        return text
    return text[: max(0, limit - 3)].rstrip() + "..."


def _trim_text_without_ellipsis(value: str, limit: int) -> str:
    text = re.sub(r"\s+", " ", str(value)).strip()
    if len(text) <= limit:
        return text
    candidate = text[:limit].rstrip()
    for marker in (".", ";", ":"):
        boundary = candidate.rfind(marker)
        if boundary >= int(limit * 0.55):
            return candidate[: boundary + 1].strip()
    space = candidate.rfind(" ")
    if space >= int(limit * 0.65):
        candidate = candidate[:space].rstrip()
    return candidate.rstrip(" ,;:")


def _row_value(row: Any, column: str, index: int) -> Any:
    if isinstance(row, dict):
        return row.get(column)
    if isinstance(row, (list, tuple)) and index < len(row):
        return row[index]
    return None


def _effective_row_count(record: dict[str, Any], query: dict[str, Any], rows: list[Any]) -> int:
    for value in (record.get("row_count"), query.get("row_count")):
        if value is None:
            continue
        try:
            return int(value)
        except (TypeError, ValueError):
            return 0
    return len(rows)


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    return _trim_text(str(value), 80)


def _string_list_from_any(values: list[Any]) -> list[str]:
    return [str(value) for value in values if str(value).strip()]


def _chunks(values: list[list[str]], size: int) -> list[list[list[str]]]:
    return [values[index:index + size] for index in range(0, len(values), size)]


def _safe_path(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(PROJECT_ROOT))
    except ValueError:
        return str(path)


def _planner_warning_codes(warnings: list[str]) -> list[str]:
    codes: list[str] = []
    for warning in warnings:
        code = str(warning).split(":", 1)[0].strip()
        if code:
            codes.append(code)
    return _unique(codes)


def _unique(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = str(value)
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return result
